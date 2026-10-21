"""Exports Excel et PDF (tableaux simples) du module Ventes & caisse."""
import io
from datetime import datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone


def _valeur(v):
    if isinstance(v, Decimal):
        return float(v) if v != v.to_integral_value() else int(v)
    return v


def construire_excel(titre, entetes, lignes, sous_titre=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = titre[:31] or "Rapport"
    ws.append([f"{titre} ({sous_titre})" if sous_titre else titre])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(list(entetes))
    for cellule in ws[3]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor="714B67")
    for ligne in lignes:
        ws.append([_valeur(v) for v in ligne])
    for i in range(1, len(entetes) + 1):
        largeur = max(len(str(entetes[i - 1])), *(len(str(l[i - 1])) for l in lignes[:200])) if lignes else len(str(entetes[i - 1]))
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, largeur + 2), 50)
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


def construire_pdf(titre, entetes, lignes, sous_titre="", entreprise=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    flux = io.BytesIO()
    doc = SimpleDocTemplate(flux, pagesize=landscape(A4), leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                            topMargin=1.2 * cm, bottomMargin=1.2 * cm)
    styles = getSampleStyleSheet()
    elements = []
    if entreprise:
        elements.append(Paragraph(entreprise, styles["Heading3"]))
    elements.append(Paragraph(titre, styles["Title"]))
    if sous_titre:
        elements.append(Paragraph(sous_titre, styles["Normal"]))
    elements.append(Spacer(1, 0.4 * cm))
    donnees = [list(entetes)] + [[str(v) if v is not None else "" for v in l] for l in lignes]
    tableau = Table(donnees, repeatRows=1)
    tableau.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#714B67")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F2F6")]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    elements.append(tableau)
    elements.append(Spacer(1, 0.3 * cm))
    elements.append(Paragraph(f"Édité le {timezone.localtime():%d/%m/%Y %H:%M}", styles["Normal"]))
    doc.build(elements)
    return flux.getvalue()


def reponse_export(fmt, slug, titre, entetes, lignes, sous_titre="", entreprise=""):
    horodatage = datetime.now().strftime("%Y%m%d_%H%M%S")
    if fmt == "xlsx":
        rep = HttpResponse(
            construire_excel(titre, entetes, lignes, sous_titre),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        rep["Content-Disposition"] = f'attachment; filename="{slug}_{horodatage}.xlsx"'
        return rep
    rep = HttpResponse(construire_pdf(titre, entetes, lignes, sous_titre, entreprise), content_type="application/pdf")
    rep["Content-Disposition"] = f'attachment; filename="{slug}_{horodatage}.pdf"'
    return rep
